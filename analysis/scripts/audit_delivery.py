"""Independent format/source audit of final vector assets and editable tables."""
from pathlib import Path
import hashlib,json,xml.etree.ElementTree as ET
from pypdf import PdfReader
from PIL import Image
from docx import Document
R=Path(__file__).resolve().parents[1]
def sha(p):return hashlib.sha256(p.read_bytes()).hexdigest()
records=[]
for folder in ['main','supplement','tables']:
    for pdf in sorted((R/'figures'/folder).glob('*.pdf')):
        reader=PdfReader(pdf);assert len(reader.pages)==1
        page=reader.pages[0];width=float(page.mediabox.width)*25.4/72
        assert abs(width-183)<.01 and len(page.images)==0
        fonts=page['/Resources']['/Font'].get_object()
        types=[str(v.get_object().get('/Subtype')) for v in fonts.values()]
        assert '/Type3' not in types
        svg=pdf.with_suffix('.svg');root=ET.fromstring(svg.read_text(encoding='utf-8'))
        assert not list(root.iter('{http://www.w3.org/2000/svg}image'))
        assert list(root.iter('{http://www.w3.org/2000/svg}text'))
        png=pdf.with_suffix('.png');im=Image.open(png)
        assert all(abs(v-1200)<.05 for v in im.info['dpi']) and im.width==8645
        records.append({'name':pdf.stem,'section':folder,'width_mm':width,'pdf_image_objects':0,
            'svg_text_editable':True,'png_pixels':list(im.size),'png_dpi':im.info['dpi'],
            'sha256':{x.suffix[1:]:sha(x) for x in [pdf,svg,png]}})
assert len(records)==14
doc=R/'delivery/Revised_River_Nutrient_Monitoring_Manuscript.docx';d=Document(doc)
expected=json.loads((R/'manuscript/table_source.json').read_text(encoding='utf-8'))
assert expected==[[[c.text for c in row.cells] for row in t.rows] for t in d.tables]
assert len(d.inline_shapes)==11 and len(d.tables)==9
frozen=json.loads((R/'figures/final_figures_manifest.json').read_text(encoding='utf-8'))
assert frozen['all_pass'] and frozen['plot_script_sha256']==sha(R/'scripts/plot_revision.py')
for rec in frozen['records']:
    actual=next(x for x in records if x['name']==rec['name'])
    assert actual['sha256']==rec['exports']['sha256']
report={'status':'PASS','figures':11,'vector_main_tables':3,'native_word_tables':9,
    'source_preserved_table_cells':sum(sum(map(len,x)) for x in expected),
    'document_sha256':sha(doc),'manuscript_sha256':sha(R/'manuscript/manuscript.md'),
    'vector_assets':records,'figure_manifest_unchanged':True,
    'note':'DPI is a raster property; PDF/SVG are genuine vector files with no raster image objects.'}
(R/'docs/DELIVERY_AUDIT.json').write_text(json.dumps(report,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in report.items() if k!='vector_assets'},indent=2))
