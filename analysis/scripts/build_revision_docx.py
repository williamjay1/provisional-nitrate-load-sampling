"""Native Word article, source-preserving tables, equations and inline figures."""
from pathlib import Path
import subprocess,json,re,zipfile
from docx import Document
from docx.shared import Mm,Pt,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree
R=Path(__file__).resolve().parents[1];TEMP=R/'temp';OUT=R/'delivery'
DEST=OUT/'Revised_River_Nutrient_Monitoring_Manuscript.docx'
def styles(d):
    for sec in d.sections:
        sec.page_width=Mm(215.9);sec.page_height=Mm(279.4)
        sec.left_margin=sec.right_margin=Mm(16.45)
        sec.top_margin=sec.bottom_margin=Mm(17)
        sec.header_distance=sec.footer_distance=Mm(8)
    for name,size in [('Normal',11),('Body Text',11),('Title',16),('Heading 1',13),('Heading 2',11.5),('Heading 3',11),('Caption',9.5)]:
        x=d.styles[name];x.font.name='Times New Roman';x.font.size=Pt(size);x.font.color.rgb=RGBColor(0,0,0)
        pf=x.paragraph_format;pf.line_spacing=1.05;pf.space_after=Pt(5);pf.widow_control=True
        if name=='Caption':x.font.bold=False
        if name.startswith('Heading'):x.font.bold=True;pf.keep_with_next=True;pf.space_before=Pt(10)
    for st in d.styles:
        for border in st.element.xpath('.//w:pBdr'):border.getparent().remove(border)
        for rf in st.element.xpath('.//w:rFonts'):
            for attr in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
                rf.attrib.pop(qn('w:'+attr),None)
def main():
    ref=TEMP/'reference_revision.docx';d=Document();styles(d);d.save(ref)
    subprocess.run(['pandoc',str(R/'manuscript/manuscript.md'),'-f','markdown-implicit_figures+tex_math_single_backslash+tex_math_dollars','--reference-doc',str(ref),'--resource-path',str(R/'manuscript'),'-o',str(DEST)],check=True)
    d=Document(DEST);styles(d);d.paragraphs[0].style=d.styles['Title']
    d.core_properties.title=d.paragraphs[0].text;d.core_properties.author='';d.core_properties.last_modified_by=''
    d.core_properties.subject='Environmental monitoring assessment; substantive revision for author review'
    for sec in d.sections:
        ln=OxmlElement('w:lnNumType');ln.set(qn('w:countBy'),'1');ln.set(qn('w:restart'),'continuous');ln.set(qn('w:distance'),'120');sec._sectPr.append(ln)
        foot=sec.footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');foot._p.append(fld)
    for p in d.paragraphs:
        for border in p._p.xpath('.//w:pBdr'):border.getparent().remove(border)
        if p.text.startswith('Table ') or p.text.startswith('Fig. ') or p.text.startswith('Figure '):
            p.style=d.styles['Caption']
            p.paragraph_format.keep_together=True
        if p.text.startswith('Table '):p.paragraph_format.keep_with_next=True
        if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True;p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        if p.text.startswith('Online Resource ') or p.text=='References':p.paragraph_format.page_break_before=True
    # Pandoc table schema can detach on this Windows LibreOffice; reconstruct native tables.
    for old in list(d.tables):
        vals=[[c.text for c in row.cells] for row in old.rows]
        new=d.add_table(rows=len(vals),cols=len(vals[0]));new.style='Normal Table'
        for row,vs in zip(new.rows,vals):
            for c,v in zip(row.cells,vs):c.text=v
        old._tbl.addprevious(new._tbl);old._tbl.getparent().remove(old._tbl)
    widths=[[43,29,29,28,24,30],[42,23.5,23.5,23.5,23.5,23.5,23.5],[22,29,36,36,30,30],[24,104,31,24],[25,14,28,19,22,36,39],[15,47,40,40,41],[15,27,30,27,42,42],[20,37,40,28,29,29],[15,27,30,37,37,37]]
    for i,t in enumerate(d.tables):
        t.autofit=False;ws=widths[i] if len(widths[i])==len(t.columns) else [183/len(t.columns)]*len(t.columns)
        ws=[x*183/sum(ws) for x in ws]
        for j,col in enumerate(t.columns):col.width=Mm(ws[j])
        pr=t._tbl.tblPr;b=OxmlElement('w:tblBorders')
        for edge in ['top','bottom','left','right','insideH','insideV']:
            x=OxmlElement('w:'+edge);x.set(qn('w:val'),'single' if edge in ['top','bottom'] else 'nil');x.set(qn('w:sz'),'8');x.set(qn('w:color'),'000000');b.append(x)
        pr.append(b);mar=OxmlElement('w:tblCellMar')
        for edge,val in [('top',70),('bottom',70),('left',75),('right',75)]:
            x=OxmlElement('w:'+edge);x.set(qn('w:w'),str(val));x.set(qn('w:type'),'dxa');mar.append(x)
        pr.append(mar)
        for ri,row in enumerate(t.rows):
            rp=row._tr.get_or_add_trPr();rp.append(OxmlElement('w:cantSplit'))
            if ri==0:rp.append(OxmlElement('w:tblHeader'))
            for j,c in enumerate(row.cells):
                c.width=Mm(ws[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                cp=c._tc.get_or_add_tcPr()
                if ri==0:
                    cb=OxmlElement('w:tcBorders');x=OxmlElement('w:bottom');x.set(qn('w:val'),'single');x.set(qn('w:sz'),'5');x.set(qn('w:color'),'000000');cb.append(x);cp.append(cb)
                for p in c.paragraphs:
                    pf=p.paragraph_format;pf.space_before=Pt(0);pf.space_after=Pt(1);pf.line_spacing=1.0;pf.keep_with_next=(i<3 and ri<len(t.rows)-1)
                    p.alignment=WD_ALIGN_PARAGRAPH.LEFT if j<2 else WD_ALIGN_PARAGRAPH.CENTER
                    for x in p.runs:x.font.name='Arial';x.font.size=Pt(9.2);x.font.bold=ri==0
    d.save(DEST)
    expected=json.loads((R/'manuscript/table_source.json').read_text(encoding='utf-8'))
    assert len(expected)==len(d.tables)
    for i,(e,t) in enumerate(zip(expected,d.tables)):
        assert e==[[c.text for c in row.cells] for row in t.rows],(i,'table text lost')
    with zipfile.ZipFile(DEST) as z:
        root=etree.fromstring(z.read('word/document.xml'));ns={'m':'http://schemas.openxmlformats.org/officeDocument/2006/math','w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
        audit={'status':'PASS','tables':len(d.tables),'native_equations':len(root.xpath('//m:oMath',namespaces=ns)),'inline_figures':len(d.inline_shapes),'line_numbering':bool(root.xpath('//w:lnNumType',namespaces=ns)),'table_cells_preserved':sum(sum(map(len,x)) for x in expected),'bytes':DEST.stat().st_size}
        assert audit['native_equations']>25 and audit['inline_figures']>=9 and audit['line_numbering']
    (OUT/'document_build_check.json').write_text(json.dumps(audit,indent=2));print(json.dumps(audit))
if __name__=='__main__':main()
