"""Convert the audited manuscript to a plain, editable Word article."""
from pathlib import Path
import json, re, subprocess, zipfile
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from lxml import etree

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'delivery';OUT.mkdir(exist_ok=True)
TEMP=ROOT/'temp';TEMP.mkdir(exist_ok=True)
SOURCE=ROOT/'manuscript/manuscript.md'
DEST=OUT/'River_Nutrient_Adaptive_Sampling_Manuscript.docx'
REF=TEMP/'article_reference.docx'

def style_document(doc):
    sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1)
    for name,size in [('Normal',11.5),('Body Text',11.5),('Title',16),('Heading 1',14),('Heading 2',12),('Heading 3',11.5),('Caption',10.5)]:
        s=doc.styles[name];s.font.name='Times New Roman';s.font.size=Pt(size);s.font.color.rgb=RGBColor(0,0,0)
        p=s.paragraph_format;p.line_spacing=1.08;p.space_after=Pt(6)
        if name.startswith('Heading'):
            s.font.bold=True;p.keep_with_next=True;p.space_before=Pt(10)
    doc.styles['Normal'].paragraph_format.widow_control=True

def main():
    d=Document();style_document(d);d.save(REF)
    subprocess.run(['pandoc',str(SOURCE),'-f','markdown-implicit_figures+tex_math_single_backslash+tex_math_dollars',
        '--reference-doc',str(REF),'--resource-path',str(ROOT/'manuscript'),'-o',str(DEST)],check=True)
    d=Document(DEST);style_document(d)
    d.core_properties.title=d.paragraphs[0].text
    d.core_properties.author='';d.core_properties.last_modified_by=''
    d.core_properties.subject='Original research article for Environmental Monitoring and Assessment'
    d.paragraphs[0].style=d.styles['Title']
    # Simple page numbering; native OMML equations produced by Pandoc are retained.
    foot=d.sections[0].footer.paragraphs[0];foot.alignment=WD_ALIGN_PARAGRAPH.CENTER
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');foot._p.append(field)
    for p in d.paragraphs:
        if p.text.startswith('Table '):p.paragraph_format.keep_with_next=True
        if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
        if p.text in ('References','Figures','Supplementary material'):
            p.paragraph_format.page_break_before=True
    widths=[
      [1.00,1.34,.42,.55,1.0,1.0,1.19],
      [1.82,.78,.78,.78,.78,.78,.78],
      [2.08,1.76,1.21,1.25],
      [.55,2.83,1.04,1.04,1.04],
      [.55,1.29,.96,.71,1.45,1.54],
      [.64,1.35,1.9,.87,.87,.87],
      [.63,.95,2.22,.91,.90,.89],
    ]
    # Rebuild table containers with python-docx's native schema. The custom
    # reference has no Pandoc Table/Compact styles, which LibreOffice can
    # otherwise interpret as detached cell text on this Windows installation.
    for old in list(d.tables):
        values=[[c.text for c in row.cells] for row in old.rows]
        new=d.add_table(rows=len(values), cols=len(values[0]))
        new.style='Table Grid'
        for row,vals in zip(new.rows,values):
            for cell,value in zip(row.cells,vals):cell.text=value
        old._tbl.addprevious(new._tbl);old._tbl.getparent().remove(old._tbl)
    for i,t in enumerate(d.tables):
        t.autofit=False
        ws=widths[i] if i<len(widths) and len(widths[i])==len(t.columns) else [6.5/len(t.columns)]*len(t.columns)
        # Normalize to the actual 6.5-inch text width.
        ws=[v*6.5/sum(ws) for v in ws]
        for j,col in enumerate(t.columns):col.width=Inches(ws[j])
        pr=t._tbl.tblPr
        borders=OxmlElement('w:tblBorders')
        for edge in ['top','left','bottom','right','insideH','insideV']:
            el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');borders.append(el)
        pr.append(borders)
        margins=OxmlElement('w:tblCellMar')
        for edge in ['top','left','bottom','right']:
            el=OxmlElement('w:'+edge);el.set(qn('w:w'),'60');el.set(qn('w:type'),'dxa');margins.append(el)
        pr.append(margins)
        for ri,row in enumerate(t.rows):
            trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
            if ri==0:trpr.append(OxmlElement('w:tblHeader'))
            for j,cell in enumerate(row.cells):
                cell.width=Inches(ws[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
                if ri==0:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F2F2F2');cell._tc.get_or_add_tcPr().append(shade)
                for p in cell.paragraphs:
                    p.paragraph_format.space_before=Pt(0);p.paragraph_format.space_after=Pt(3)
                    p.paragraph_format.line_spacing=1.0;p.paragraph_format.keep_with_next=False
                    p.alignment=WD_ALIGN_PARAGRAPH.CENTER if j>=len(ws)-3 else WD_ALIGN_PARAGRAPH.LEFT
                    for r in p.runs:r.font.name='Times New Roman';r.font.size=Pt(10);r.font.bold=(ri==0)
    d.save(DEST)
    blocks=[]
    for match in re.finditer(r'(?m)^\|[^\n]*\n\|[-:| ]+\n(?:\|[^\n]*\n)+',SOURCE.read_text(encoding='utf-8')):
        rows=[line.strip().strip('|').split('|') for line in match.group().strip().splitlines()]
        blocks.append([[v.strip() for v in row] for row in rows[:1]+rows[2:]])
    assert len(blocks)==len(d.tables)==7
    table_checks=[]
    for i,(expected,t) in enumerate(zip(blocks,d.tables),1):
        actual=[[c.text for c in row.cells] for row in t.rows]
        assert expected==actual,(i,'table source mismatch')
        table_checks.append({'table':i,'rows_including_header':len(actual),'cells':sum(map(len,actual)),'status':'PASS'})
    (OUT/'document_table_check.json').write_text(json.dumps({'status':'PASS','table_checks':table_checks,
        'all_table_cells_preserved':sum(c['cells'] for c in table_checks)},indent=2),encoding='utf-8')
    with zipfile.ZipFile(DEST) as z:
        xml=etree.fromstring(z.read('word/document.xml'))
        ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main','m':'http://schemas.openxmlformats.org/officeDocument/2006/math'}
        report={'docx':str(DEST),'bytes':DEST.stat().st_size,'paragraphs':len(d.paragraphs),'tables':len(d.tables),
            'inline_shapes':len(d.inline_shapes),'native_equations':len(xml.xpath('//m:oMath',namespaces=ns)),
            'source_words':len(SOURCE.read_text(encoding='utf-8').split())}
        assert report['tables']==7 and report['inline_shapes']==3 and report['native_equations']>25,report
    (OUT/'document_build_check.json').write_text(json.dumps(report,indent=2),encoding='utf-8');print(json.dumps(report))

if __name__=='__main__':main()
